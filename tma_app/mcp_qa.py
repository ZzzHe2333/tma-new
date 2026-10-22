"""Real MCP client round-trip checks, also runnable inside the frozen build."""
import asyncio
import csv
import hashlib
import json
import os
from pathlib import Path
import socket
import time
import traceback
import uuid

from . import __version__
from . import mcp_service as mcp


def free_port():
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def wait_running(service):
    deadline = time.monotonic() + 25
    while service.state == "starting" and time.monotonic() < deadline:
        time.sleep(.025)
    if service.state != "running":
        raise AssertionError(f"MCP 启动失败：{service.state} {service.error}")


async def transport_checks(settings, files, directory):
    import httpx
    from mcp import ClientSession
    from mcp.client.streamable_http import streamable_http_client
    from openpyxl import load_workbook
    from PIL import Image
    from .batch_analysis import calculate_file

    checks = []
    before = {str(path): hashlib.sha256(Path(path).read_bytes()).hexdigest() for path in files}
    url = mcp.endpoint(settings["port"])
    headers = {"Authorization": "Bearer " + settings["token"]}
    async with httpx.AsyncClient(trust_env=False, timeout=20) as client:
        assert (await client.post(url, json={})).status_code == 401
        assert (await client.post(url, json={}, headers={"Authorization": "Bearer wrong"})).status_code == 401
        assert (await client.post(url, json={}, headers={**headers, "Origin": "https://untrusted.example"})).status_code == 403
        assert (await client.post(url, json={}, headers={**headers, "Host": "untrusted.example"})).status_code == 421
    checks.append("bearer_token_host_and_origin_enforced")
    async with httpx.AsyncClient(headers=headers, trust_env=False, timeout=30) as client:
        async with streamable_http_client(url, http_client=client) as (reader, writer, _):
            async with ClientSession(reader, writer) as session:
                await session.initialize()
                listed = await session.list_tools()
                names = {tool.name for tool in listed.tools}
                assert names == {"tma_status", "tma_batch_analyze", "tma_batch_export", "tma_job_status", "tma_cancel_job"}
                checks.append("official_mcp_initialize_list_tools")

                async def call(name, args=None):
                    result = await session.call_tool(name, args or {})
                    if result.isError:
                        raise AssertionError(str(result.content))
                    return result.structuredContent or json.loads(result.content[0].text)

                async def finished(key):
                    deadline = time.monotonic() + 45
                    while time.monotonic() < deadline:
                        value = await call("tma_job_status", {"job_id": key})
                        if value["state"] in ("completed", "failed", "cancelled"):
                            return value
                        await asyncio.sleep(.1)
                    raise AssertionError("批量任务超时")

                status = await call("tma_status")
                assert status["version"] == __version__ and status["algorithm"] == "V233"
                assert "token" not in status
                bad = await session.call_tool("tma_batch_analyze", {"files": ["relative.001"]})
                assert bad.isError
                checks.append("typed_tool_errors_and_status")
                request = [str(Path(files[0]).resolve()), str(Path(files[-1]).resolve()), str(Path(files[0]).resolve())]
                submitted = await call("tma_batch_export", {"files": request,
                    "output_directory": str(Path(directory).resolve()), "formats": ["png", "xlsx", "csv"]})
                job = await finished(submitted["job_id"])
                assert job["state"] == "completed", job
                assert [row["file"] for row in job["rows"]] == request
                assert [row["M"] for row in job["rows"]] == [.47, .48, .47]
                assert all(not row["error"] for row in job["rows"])
                assert len(job["artifacts"]) == 5
                paths = [item["path"] for item in job["artifacts"]]
                assert len(set(paths)) == len(paths) and all(Path(path).is_file() for path in paths)
                for row, file in zip(job["rows"], request):
                    reference = calculate_file(file)
                    assert all(row[key] == reference[key] for key in ("T1", "T2", "H"))
                checks.append("batch_v233_order_duplicate_identity_mass_match")
                xlsx = next(item["path"] for item in job["artifacts"] if item["type"] == "xlsx")
                book = load_workbook(xlsx, data_only=True)
                sheet = book.active
                assert sheet.max_column == 4 and sheet.max_row == 4
                for index, row in enumerate(job["rows"], 2):
                    for column, key in enumerate(("T1", "T2", "M", "H"), 1):
                        assert abs(sheet.cell(index, column).value - row[key]) < 1e-8
                book.close()
                csv_path = next(item["path"] for item in job["artifacts"] if item["type"] == "csv")
                with Path(csv_path).open(encoding="utf-8-sig", newline="") as handle:
                    cells = list(csv.reader(handle))
                assert len(cells) == 4 and len(cells[0]) == 4
                assert float(cells[1][2]) == .47
                assert len(job["tsv"].splitlines()) == 4
                for artifact in job["artifacts"]:
                    if artifact["type"] == "png":
                        with Image.open(artifact["path"]) as image:
                            assert image.width > 1920 and image.height > 1080
                            assert image.info.get("Author") == "Power By ZzzHe"
                checks.append("png_white_border_credit_xlsx_csv_tsv_numeric_readback")
                missing = str(Path(directory).resolve() / "missing.001")
                submitted = await call("tma_batch_analyze", {"files": [request[0], missing], "mass_source": ""})
                partial = await finished(submitted["job_id"])
                assert partial["rows"][0]["M"] is None
                assert partial["rows"][1]["error"] and partial["rows"][1]["T1"] is None
                assert partial["processed"] == 2
                checks.append("missing_file_and_mass_reported_without_fabrication")
                cancelled = await call("tma_batch_analyze", {"files": request * 20})
                await call("tma_cancel_job", {"job_id": cancelled["job_id"]})
                cancelled = await finished(cancelled["job_id"])
                assert cancelled["state"] == "cancelled"
                checks.append("cancel_job_via_mcp")
    assert before == {str(path): hashlib.sha256(Path(path).read_bytes()).hexdigest() for path in files}
    checks.append("input_hashes_unchanged")
    return checks, job


def verify_mcp(window, output, files):
    output.parent.mkdir(parents=True, exist_ok=True)
    report = {"status": "failed", "version": __version__, "checks": []}
    original_appdata, original_mass = os.environ.get("LOCALAPPDATA"), os.environ.get("TMA_MASS_XLSX")
    os.environ["LOCALAPPDATA"] = str(output.parent / ("mcp-config-" + uuid.uuid4().hex[:8]))
    try:
        assert len(files) >= 2
        assert not mcp.load_settings()["enabled"] and window.mcp_service.state == "stopped"
        report["checks"].append("fresh_config_disabled_no_listener")
        sheet = output.with_suffix(".weights.csv")
        with sheet.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle); writer.writerow(["样品名", "M (mg)"])
            for index, path in enumerate(dict.fromkeys(files)):
                writer.writerow([path.stem, .47 + index*.01])
        os.environ["TMA_MASS_XLSX"] = str(sheet.resolve())
        settings = {"enabled": True, "port": free_port(), "token": mcp.new_token()}
        mcp.save_settings(settings)
        window.start_saved_mcp()
        wait_running(window.mcp_service)
        checks, job = asyncio.run(transport_checks(settings, files, output.parent / "exports"))
        report["checks"].extend(checks)
        report["artifacts"] = job["artifacts"]
        # Use the real settings UI to disable, then verify persistence and socket closure.
        window.show_mcp_settings()
        dialog = window.mcp_dialog
        dialog.enabled.setChecked(False)
        dialog.apply()
        assert not mcp.load_settings()["enabled"]
        assert window.mcp_service.state == "stopped"
        with socket.socket() as probe:
            assert probe.connect_ex(("127.0.0.1", settings["port"])) != 0
        dialog.close()
        report["checks"].append("settings_ui_disable_persist_and_release_port")
        report["status"] = "passed"
    except BaseException:
        report["error"] = traceback.format_exc()
    finally:
        window.mcp_service.stop()
        for key, value in (("LOCALAPPDATA", original_appdata), ("TMA_MASS_XLSX", original_mass)):
            if value is None: os.environ.pop(key, None)
            else: os.environ[key] = value
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        window.dirty = False
        window.close()
    return report
