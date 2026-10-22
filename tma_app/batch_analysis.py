"""V14/V233 batch calculations and exports, independent of either GUI."""
from __future__ import annotations

import csv
import hashlib
from io import StringIO
import logging
import math
import os
import json
from pathlib import Path
import tempfile
import threading
from functools import wraps

import numpy as np
from . import piliang_core as original
from .mass_lookup import lookup_record

KEYS = ("T1", "T2", "M", "H")
HEADERS = ("T1 (°C)", "T2 (°C)", "M (mg)", "H (µm)")
LOGGER = logging.getLogger(__name__)
_BATCH_LOCK = threading.RLock()


def serialized(function):
    """GUI batch jobs and MCP share algorithm/cache/Agg export operations."""
    @wraps(function)
    def wrapped(*args, **kwargs):
        with _BATCH_LOCK:
            return function(*args, **kwargs)
    return wrapped


def history_path():
    from .path_settings import config_path
    return config_path().parent / "batch_history.json"


def history_snapshot(rows):
    """Keep metrics/provenance, not raw arrays or temporary export errors."""
    keys = {"id", "batch_id", "batch_label", "imported_at", "file", "name", "T1", "T2", "M", "H",
            "H1", "H2", "error", "warning", "mass_source", "mass_config", "state", "png", "sha256",
            "algorithm", "algorithm_source_sha256", "reader"}
    return [{key: value for key, value in row.items() if key in keys} for row in rows]


def load_history():
    path = history_path()
    if not path.exists():
        return []
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    rows = data.get("rows") if isinstance(data, dict) else None
    if not isinstance(rows, list):
        raise ValueError("批量历史记录格式无效")
    seen = set()
    for row in rows:
        if not isinstance(row, dict) or any(not isinstance(row.get(k), str) for k in ("id", "batch_id", "batch_label", "file", "name", "state")) or row["id"] in seen:
            raise ValueError("批量历史记录内容无效")
        seen.add(row["id"])
        if any(row.get(key) is not None and (type(row[key]) not in (int, float) or not math.isfinite(row[key])) for key in (*KEYS, "H1", "H2")):
            raise ValueError("批量历史数值无效")
        if row["state"] in ("计算中", "排队中", "匹配 M…"):
            row["state"] = "未完成"
    return rows


def save_history(rows):
    path = history_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix="batch-history-", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump({"version": 1, "rows": history_snapshot(rows)}, handle, ensure_ascii=False, allow_nan=False)
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return str(path)


@serialized
def calculate_file(path):
    """Use exactly the supplied reader, V233 detection and raw interpolation."""
    result = dict(file=str(path), name=Path(path).name, T1=None, T2=None, M=None, H=None,
                  error="", warning="", mass_source="", state="计算中", png="")
    try:
        from .cloud_storage import resolve_source
        path = resolve_source(path)
        result.update(file=str(path), name=path.name)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        text, read_meta = original.load_text_prefer_tajm(str(path), LOGGER)
        xs, ys, _ = original.extract_xy_from_file(original.parse_lines(text), LOGGER)
        if len(xs) < 5:
            raise ValueError(f"有效数据点不足（{len(xs)} 行）")
        t1, t2, details = original.detect_T2_T1_super_gradient(xs, ys, LOGGER)
        cache = original.build_curve_cache(xs, ys)
        h1, h2 = (original.interpolate_curve_height(t, cache) for t in (t1, t2))
        if not np.isfinite([t1, t2, h1, h2]).all():
            raise ValueError("算法返回了非有限数值")
        result.update(T1=t1, T2=t2, H=h2-h1, H1=h1, H2=h2, xs=xs, ys=ys,
                      sha256=digest, details=details, algorithm=original.ALGORITHM_VERSION,
                      algorithm_source_sha256=original.SOURCE_SHA256, reader=read_meta.get("used"),
                      state="匹配 M…")
    except Exception as exc:
        result.update(error=str(exc), state="读取失败")
    return result


@serialized
def match_mass(result, source):
    """Retain file-name matching; configured sources support mg or explicit g."""
    result = dict(result)
    if result.get("error"):
        return result
    match = lookup_record(result["name"], source=source)
    mass = match.value
    if mass is not None:
        mass *= 1000 if match.unit == "g" else 1
    if mass is None or not math.isfinite(mass) or mass <= 0:
        result.update(M=None, warning=match.message or "未匹配有效重量", state="M 未匹配")
    else:
        result.update(M=mass, warning="", state="完成")
    result.update(mass_source=match.source, mass_config=source)
    return result


def display_value(row, key):
    value = row.get(key)
    if value is None:
        return ""
    if key in ("T1", "T2"):
        return original.format_plot_temp_c(value)
    if key == "H":
        return original.format_height_um(value)
    return f"{value:.4f}"


def table_text(rows, include_headers=True):
    """Four tab-separated columns, directly pasteable into Excel/WPS."""
    buffer = StringIO(newline="")
    writer = csv.writer(buffer, delimiter="\t", lineterminator="\r\n")
    if include_headers:
        writer.writerow(HEADERS)
    for row in rows:
        writer.writerow([display_value(row, key) for key in KEYS])
    return buffer.getvalue()


def protect_output(path, rows):
    target = os.path.normcase(str(Path(path).resolve()))
    for row in rows:
        for source in (row.get("file"), row.get("mass_config"), row.get("mass_source", "").split(" [", 1)[0]):
            if source and os.path.normcase(str(Path(source).resolve())) == target:
                raise ValueError("输出不能覆盖原始数据或称量表")


@serialized
def write_xlsx(rows, path):
    """Four numeric columns; file identity/status is stored in cell comments."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    from .cloud_storage import touch_path
    for row in rows: touch_path(row["file"])
    path = Path(path)
    protect_output(path, rows)
    workbook = Workbook()
    workbook.properties.creator = "ZzzHe"
    workbook.properties.title = "T1 T2 M H · V233"
    sheet = workbook.active
    sheet.title = "批量结果"
    sheet.append(HEADERS)
    for index, row in enumerate(rows, 2):
        sheet.append([row.get(key) for key in KEYS])
        note = f"文件：{row['name']}\n路径：{row['file']}\n算法：V233\n状态：{row['state']}"
        if row.get("mass_source"):
            note += f"\nM 来源：{row['mass_source']}"
        if row.get("error") or row.get("warning"):
            note += "\n" + (row.get("error") or row.get("warning"))
        sheet.cell(index, 1).comment = Comment(note, "ZzzHe")
        for column, key in enumerate(KEYS, 1):
            cell = sheet.cell(index, column)
            cell.alignment = Alignment(horizontal="right")
            if key == "M":
                cell.number_format = "0.0000"
            elif key == "H":
                cell.number_format = "0" if abs(row.get(key) or 0) > 1000 else "0.0"
            else:
                cell.number_format = "0.00"
            if row.get(key) is None:
                cell.fill = PatternFill("solid", fgColor="FFF1E4")
        sheet.row_dimensions[index].height = 23
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="674B7B")
        cell.fill = PatternFill("solid", fgColor="EEE3F7")
    for column in "ABCD":
        sheet.column_dimensions[column].width = 23
    sheet.row_dimensions[1].height = 26
    sheet.freeze_panes = "A2"
    if rows:
        table = Table(displayName="TMAResults", ref=f"A1:D{len(rows)+1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight20", showRowStripes=True)
        sheet.add_table(table)
    descriptor, temporary = tempfile.mkstemp(prefix="tma-batch-", suffix=".xlsx", dir=path.parent)
    os.close(descriptor)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        Path(temporary).unlink(missing_ok=True)
    return str(path.resolve())


def plot_paths(rows, directory):
    """Reserve unique names in input order, including same-name/same-file runs."""
    reserved = set()
    result = []
    for row in rows:
        stem = Path(row["file"]).stem
        candidate = Path(directory) / (stem + ".png")
        index = 2
        while os.path.normcase(str(candidate)) in reserved or candidate.exists():
            candidate = Path(directory) / f"{stem} ({index}).png"
            index += 1
        reserved.add(os.path.normcase(str(candidate)))
        result.append(candidate)
    return result


@serialized
def write_plot(row, path):
    """V14 cross/temperature/height layout, with the app's white border and credit."""
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.ticker import MultipleLocator, AutoMinorLocator

    from .cloud_storage import resolve_source
    source_file = resolve_source(row["file"])
    if row.get("error") or row.get("T1") is None:
        raise ValueError("此文件没有可出图的结果")
    if "xs" not in row or "ys" not in row:
        row = dict(row)
        source = source_file
        if hashlib.sha256(source.read_bytes()).hexdigest() != row.get("sha256"):
            raise ValueError("原始文件已变化，请重新导入计算后出图")
        text, _ = original.load_text_prefer_tajm(str(source), LOGGER)
        row["xs"], row["ys"], _ = original.extract_xy_from_file(original.parse_lines(text), LOGGER)
    path = Path(path)
    protect_output(path, [row])
    figure = Figure(figsize=(12.8, 7.2), dpi=150, facecolor="white")
    FigureCanvasAgg(figure)
    axis = figure.add_subplot(111)
    axis.plot(row["xs"], row["ys"], color="black", linewidth=.85)
    t1, t2, h1, h2 = (row[key] for key in ("T1", "T2", "H1", "H2"))
    axis.plot([t1, t2], [h1, h2], "+", color="black", ms=15, mew=.9)
    axis.annotate(f"{t1:.2f} °C", (t1, h1), xytext=(0, -20), textcoords="offset points", ha="center")
    axis.annotate(f"{t2:.2f} °C", (t2, h2), xytext=(0, 15), textcoords="offset points", ha="center")
    axis.plot([t1, t2], [h1, h1], color="black", linewidth=.65)
    axis.annotate("", (t2, h2), (t2, h1), arrowprops={"arrowstyle": "<->", "color": "black", "lw": .8})
    axis.annotate(f"{original.format_height_um(abs(row['H']))} µm", (t2, (h1+h2)/2), xytext=(12, 0), textcoords="offset points")
    lo, hi = min(row["ys"]), max(row["ys"])
    span = max(hi-lo, 1)
    axis.set_ylim(lo-span*.15, hi+span*.2)
    axis.set_xlim(min(0., min(row["xs"])), max(300., max(row["xs"])))
    axis.xaxis.set_major_locator(MultipleLocator(50))
    axis.xaxis.set_minor_locator(MultipleLocator(10))
    axis.yaxis.set_minor_locator(AutoMinorLocator())
    axis.grid(which="major", color="#e6e0e9", linestyle="--", lw=.4)
    axis.set_xlabel("温度(°C)")
    axis.set_ylabel("尺寸变化(µm)")
    figure.text(.085, .965, row["name"], ha="left", va="top", fontsize=12, parse_math=False)
    mass = "未匹配" if row.get("M") is None else f"{row['M']:.2f} mg"
    axis.text(.02, .96, "方法日志：30 ℃ 平衡，施加 0.06 N\n                    20 ℃/min加热至 250 ℃\n称取质量：" + mass,
              transform=axis.transAxes, va="top", fontsize=10)
    figure.text(.97, .025, "Power By ZzzHe", ha="right", va="bottom", fontsize=10)
    figure.subplots_adjust(left=.085, right=.96, top=.9, bottom=.12)
    descriptor, temporary = tempfile.mkstemp(prefix="tma-plot-", suffix=".png", dir=path.parent)
    os.close(descriptor)
    try:
        figure.savefig(temporary, format="png", dpi=150, facecolor="white", bbox_inches=figure.bbox_inches.padded(.25),
                       metadata={"Software": "ZzzHe TMA Analysis", "Author": "Power By ZzzHe"})
        os.replace(temporary, path)
    finally:
        figure.clear()
        Path(temporary).unlink(missing_ok=True)
    return str(path.resolve())
