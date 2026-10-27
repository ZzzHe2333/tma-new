"""Exercise the packaged batch UI against real inputs in an isolated config."""
import csv
import hashlib
import json
from pathlib import Path
import os
import traceback
import uuid
import threading

from PySide6.QtCore import QEventLoop, QTimer
from PySide6.QtWidgets import QApplication, QFileDialog

from . import __version__
from . import batch_analysis as backend


def wait_until(callback, timeout=30000):
    if callback():
        return
    loop = QEventLoop()
    timer = QTimer()
    timer.timeout.connect(lambda: loop.quit() if callback() else None)
    timer.start(10)
    deadline = QTimer()
    deadline.setSingleShot(True)
    deadline.timeout.connect(loop.quit)
    deadline.start(timeout)
    loop.exec()
    timer.stop()
    deadline.stop()
    if not callback():
        raise ValueError("批量操作验证超时")


def verify_batch(window, output, paths):
    output.parent.mkdir(parents=True, exist_ok=True)
    report = {"status": "failed", "version": __version__, "checks": []}
    original_appdata = os.environ.get("LOCALAPPDATA")
    original_mass = os.environ.get("TMA_MASS_XLSX")
    # Even a manual invocation cannot write the real user's settings/history.
    os.environ["LOCALAPPDATA"] = str(output.parent / (output.stem + "-config-" + uuid.uuid4().hex[:8]))
    try:
        if len(paths) < 2:
            raise ValueError("验证至少需要两个 .001 文件")
        before = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
        sheet = output.with_suffix(".weights.csv")
        with sheet.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["样品名", "M (mg)"])
            for index, path in enumerate(dict.fromkeys(paths)):
                writer.writerow([path.stem, .47 + index*.01])
        os.environ["TMA_MASS_XLSX"] = str(sheet)
        window.import_files(paths[:2])
        dialog = window.batch_dialog
        assert dialog.isVisible()
        wait_until(lambda: dialog.job is None and window.worker is None)
        assert len(dialog.rows) == 2 and all(row.get("M") for row in dialog.rows)
        first_batch = dialog.latest_batch_id
        report["checks"].append("normal_multi_import_auto_calculates_and_matches_mass")
        dialog.enqueue_files([paths[0]])
        wait_until(lambda: dialog.job is None)
        assert len(dialog.rows) == 3 and len(dialog.view_rows) == 1
        dialog.grab().save(str(output.with_suffix(".latest.png")))
        dialog.all_button.click()
        assert len(dialog.view_rows) == 3
        dialog.grab().save(str(output.with_suffix(".all.png")))
        dialog.batch_filter.setCurrentIndex(dialog.batch_filter.findData(first_batch))
        assert len(dialog.view_rows) == 2
        dialog.CopyBatch.click()
        clipboard = QApplication.clipboard().text()
        assert len(clipboard.splitlines()) == 3
        assert all(len(line.split("\t")) == 4 for line in clipboard.splitlines())
        output.with_suffix(".clipboard.tsv").write_text(clipboard, encoding="utf-8-sig")
        report["checks"].append("latest_all_batch_filter_and_four_column_clipboard")
        save_picker = QFileDialog.getSaveFileName
        folder_picker = QFileDialog.getExistingDirectory
        xlsx = output.with_suffix(".xlsx")
        plot_folder = output.parent / (output.stem + "-plots")
        plot_folder.mkdir(exist_ok=True)
        try:
            QFileDialog.getSaveFileName = staticmethod(lambda *a, **kw: (str(xlsx), ""))
            QFileDialog.getExistingDirectory = staticmethod(lambda *a, **kw: str(plot_folder))
            dialog.ExportBatchXlsx.click()
            wait_until(lambda: dialog.export_job is None)
            from openpyxl import load_workbook
            workbook = load_workbook(xlsx, data_only=True)
            sheet_out = workbook.active
            assert sheet_out.max_column == 4 and sheet_out.max_row == 3
            for index, row in enumerate(dialog.view_rows, 2):
                for column, key in enumerate(backend.KEYS, 1):
                    assert abs(sheet_out.cell(index, column).value - row[key]) < 1e-8
                assert row["file"] in sheet_out.cell(index, 1).comment.text
            workbook.close()
            dialog.ExportBatchPlots.click()
            wait_until(lambda: dialog.export_job is None)
            from PIL import Image
            for row in dialog.view_rows:
                assert row.get("png") and not row.get("export_error")
                with Image.open(row["png"]) as picture:
                    assert picture.info["Author"] == "Power By ZzzHe" and picture.width >= 1994
                    assert picture.convert("RGB").getpixel((0, 0)) == (255, 255, 255)
            report["checks"].append("bottom_buttons_export_numeric_xlsx_and_white_border_pngs")
        finally:
            QFileDialog.getSaveFileName = save_picker
            QFileDialog.getExistingDirectory = folder_picker
        dialog.close()
        settings = window.show_path_settings()
        settings.batch_popup.setChecked(False)
        settings.grab().save(str(output.with_suffix(".settings.png")))
        settings.accept()
        wait_until(lambda:not settings.saving)
        from .path_settings import batch_popup_enabled
        assert not batch_popup_enabled()
        window.import_files(paths[:2])
        assert not dialog.isVisible()
        wait_until(lambda: dialog.job is None and window.worker is None)
        assert len(dialog.rows) == 5 and len(dialog.view_rows) == 2
        assert window.show_batch_analysis() is dialog and dialog.isVisible()
        report["checks"].append("popup_disabled_still_computes_and_menu_reopens")
        # A weighing share can stall after T1/T2/H are already calculated.
        # Exercise the actual buttons before allowing that lookup to return.
        from .mass_lookup import MassMatch
        entered,release=threading.Event(),threading.Event()
        previous_lookup=backend.lookup_record
        def slow_mass(*args,**kwargs):
            entered.set();release.wait(45)
            return MassMatch(.47,'QA controlled source','mg','matched')
        pending_xlsx=output.with_suffix('.pending.xlsx')
        pending_pngs=output.parent/(output.stem+'-pending-plots');pending_pngs.mkdir(exist_ok=True)
        try:
            backend.lookup_record=slow_mass
            dialog.enqueue_files([paths[0]])
            wait_until(lambda:entered.is_set() and dialog.view_rows[0].get('T1') is not None)
            assert dialog.job is not None and all(button.isEnabled() for button in (dialog.CopyBatch,dialog.ExportBatchXlsx,dialog.ExportBatchPlots))
            dialog.grab().save(str(output.with_suffix('.waiting-m.png')))
            dialog.CopyBatch.click()
            assert QApplication.clipboard().text().splitlines()[1].split('\t')[2]==''
            QFileDialog.getSaveFileName=staticmethod(lambda *a,**kw:(str(pending_xlsx),''))
            QFileDialog.getExistingDirectory=staticmethod(lambda *a,**kw:str(pending_pngs))
            dialog.ExportBatchXlsx.click();wait_until(lambda:dialog.export_job is None)
            book=load_workbook(pending_xlsx,data_only=True)
            assert book.active.max_row==2 and book.active.cell(2,3).value is None
            book.close()
            dialog.ExportBatchPlots.click();wait_until(lambda:dialog.export_job is None)
            pending_png=dialog.view_rows[0]['png'];assert Path(pending_png).is_file()
            assert dialog.job is not None and not release.is_set()
        finally:
            release.set();backend.lookup_record=previous_lookup
            QFileDialog.getSaveFileName=save_picker
            QFileDialog.getExistingDirectory=folder_picker
        wait_until(lambda:dialog.job is None)
        assert dialog.view_rows[0]['M']==.47 and dialog.view_rows[0]['png']==pending_png
        report['checks'].append('copy_xlsx_png_available_while_mass_is_blocked_and_late_mass_keeps_export')
        from .batch_dialog import BatchDialog
        restored = BatchDialog(window)
        assert len(restored.rows) == 6 and len(restored.view_rows) == 1
        restored.all_button.click()
        assert len(restored.view_rows) == 6
        restored.close()
        from .classic import ClassicMainWindow
        fresh=ClassicMainWindow()
        try:assert fresh.show_batch_analysis().rows==[]
        finally:fresh.dirty=False;fresh.close()
        from .path_settings import config_path
        assert not (config_path().parent/'batch_history.json').exists()
        report["checks"].append("same_app_reopens_batches_new_app_empty_and_no_disk_history")
        assert all(hashlib.sha256(p.read_bytes()).hexdigest() == before[str(p)] for p in paths)
        report["checks"].append("source_inputs_unchanged")
        report.update(status="passed", rows=len(dialog.rows), batches=len({row["batch_id"] for row in dialog.rows}),
                      algorithm="V233", algorithm_source_sha256=backend.original.SOURCE_SHA256)
    except Exception as exc:
        report.update(error=str(exc), traceback=traceback.format_exc())
        raise
    finally:
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        # Close before restoring the environment; the application clears RAM rows.
        if window.worker is None:
            window.dirty = False
            window.close()
        window.batch_dialog = None
        if original_appdata is None:
            os.environ.pop("LOCALAPPDATA", None)
        else:
            os.environ["LOCALAPPDATA"] = original_appdata
        if original_mass is None:
            os.environ.pop("TMA_MASS_XLSX", None)
        else:
            os.environ["TMA_MASS_XLSX"] = original_mass
