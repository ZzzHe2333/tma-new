from copy import deepcopy
import numpy as np
import pytest
from PySide6.QtWidgets import QApplication,QFileDialog
from openpyxl import load_workbook

from test_peak_integration import ui,data,entry
from tma_app.area_calculation import statistics,calculate,validate_snapshot
from tma_app.workspace import Curve,save_session,load_session


def test_zero_crossing_signed_and_absolute_areas_are_exact():
    s=statistics([0,2],[-2,2])
    assert s==dict(net_area=0.,positive_area=1.,negative_area=-1.,absolute_area=2.)
    # The incorrect trapezoid-of-absolute-samples formula would return 4.
    assert statistics([0,2],[2,-2])==s
    s=statistics([2,0],[2,2]);assert s['net_area']==-4 and s['positive_area']==4 and s['absolute_area']==4
    assert statistics([0,0,2],[100,2,2])['net_area']==4
    with pytest.raises(ValueError):statistics([0,2,1],[0,2,0])


@pytest.mark.parametrize('baseline',['linear','zero','left','right'])
def test_independent_baselines_and_fractional_boundaries_match_trapezoids(tmp_path,baseline):
    from scipy.integrate import trapezoid
    curve=Curve(data(tmp_path));before=curve.dataset.values.copy()
    v=calculate(curve,'temperature',[10.5,90.25],coordinate='recorded_time',baseline=baseline)
    p=v['area_profile'];r=np.array(p['signal'])-p['baseline']
    assert v['integral']==pytest.approx(trapezoid(r,p['q']))
    assert v['net_area']==pytest.approx(v['positive_area']+v['negative_area'])
    assert v['absolute_area']==pytest.approx(v['positive_area']-v['negative_area'])
    validate_snapshot(p,v,len(curve.active_rows));np.testing.assert_array_equal(curve.dataset.values,before)


def test_time_and_temperature_areas_differ_and_reversals_or_gaps_are_explicit(tmp_path):
    curve=Curve(data(tmp_path))
    a=calculate(curve,'temperature',[0,100]);b=calculate(curve,'temperature',[0,100],coordinate='display_x')
    assert a['integral']==pytest.approx(100) and b['integral']==pytest.approx(1000)
    curve.dataset.values[50,1]=curve.dataset.values[49,1]-1
    with pytest.raises(ValueError,match='回摆'):calculate(curve,'temperature',[0,100],coordinate='display_x')
    assert np.isfinite(calculate(curve,'temperature',[0,100])['integral'])
    curve.data_indices=list(range(20))+list(range(30,101))
    with pytest.raises(ValueError,match='排除'):calculate(curve,'temperature',[0,90])


def test_existing_peak_area_gains_auditable_components_without_changing_integral(tmp_path):
    from tma_app.peak_integration import values
    v=values(Curve(data(tmp_path)),entry(),[0,100])
    assert v['integral']==pytest.approx(100) and v['absolute_area']==pytest.approx(100)
    validate_snapshot(v['area_profile'],v)
    changed=deepcopy(v);changed['absolute_area']+=1
    with pytest.raises(ValueError,match='来源快照'):validate_snapshot(changed['area_profile'],changed)


def test_ui_compute_copy_export_and_session_preserve_source_and_active_document(ui,tmp_path,monkeypatch):
    ui.finish_import([data(tmp_path)],[]);document=ui.active_document();curve=ui.selected();before=curve.dataset.values.copy()
    assert ui.classic_actions['RunningIntegral1'].isEnabled()
    dialog=ui.show_area_calculation();dialog.compute()
    assert len(ui.results)==1 and ui.results[0]['values']['integral']==pytest.approx(100),dialog.feedback.text()
    assert ui.active_document() is document and ui.selected() is curve
    assert len(ui._result_annotations)==2
    dialog.copy();assert '净面积' in QApplication.clipboard().text() and '100' in QApplication.clipboard().text()
    path=tmp_path/'areas.xlsx';monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *args:(str(path),''));dialog.export()
    book=load_workbook(path,data_only=True);assert book.active.cell(2,6).value==pytest.approx(100);assert '公式与方法' in book.sheetnames;book.close()
    path=tmp_path/'area.tmas';save_session(path,ui.curves,ui.domain,ui.results,ui.plot_options)
    assert load_session(path)[2]==ui.results
    ui.write_plot(tmp_path/'area.png');np.testing.assert_array_equal(curve.dataset.values,before)
    dialog.reject()
