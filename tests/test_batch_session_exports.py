"""Session-only batches and export controls independent of slow M matching."""
import json
from pathlib import Path
import threading
import pytest

from PySide6.QtWidgets import QApplication,QFileDialog
from openpyxl import load_workbook

from test_batch_analysis import ui,samples,wait_until
from tma_app import batch_analysis as backend
from tma_app.mass_lookup import MassMatch
from tma_app.path_settings import config_path


def test_copy_and_exports_work_while_mass_lookup_is_waiting(ui,samples,tmp_path,monkeypatch):
    entered,release=threading.Event(),threading.Event()
    def slow_match(*args,**kwargs):
        entered.set();release.wait(15)
        return MassMatch(.47,'test.csv','mg','matched')
    monkeypatch.setattr(backend,'lookup_record',slow_match)
    dialog=ui.show_batch_analysis([samples[0]])
    target=tmp_path/'during-match.xlsx'
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a,**k:(str(target),''))
    monkeypatch.setattr(QFileDialog,'getExistingDirectory',lambda *a,**k:str(tmp_path))
    try:
        wait_until(lambda:entered.is_set() and dialog.rows[0].get('T1') is not None)
        assert dialog.job is not None
        assert all(b.isEnabled() for b in [dialog.ExportBatchPlots,dialog.ExportBatchXlsx,dialog.CopyBatch])
        dialog.CopyBatch.click()
        lines=QApplication.clipboard().text().splitlines()
        assert len(lines)==2 and lines[1].split('\t')[2]==''
        dialog.ExportBatchXlsx.click()
        wait_until(lambda:target.is_file() and dialog.export_job is None)
        workbook=load_workbook(target,data_only=True)
        assert workbook.active.cell(2,1).value==pytest.approx(dialog.rows[0]['T1'])
        assert workbook.active.cell(2,3).value is None
        workbook.close()
        dialog.ExportBatchPlots.click()
        wait_until(lambda:bool(dialog.rows[0].get('png')) and dialog.export_job is None)
        png=dialog.rows[0]['png'];assert Path(png).is_file()
        assert not release.is_set() and dialog.job is not None
    finally:
        release.set()
    wait_until(lambda:dialog.job is None)
    assert dialog.rows[0]['M']==.47 and dialog.rows[0]['png']==png


def test_old_disk_history_is_ignored_and_never_updated(ui,samples):
    legacy=config_path().parent/'batch_history.json'
    legacy.parent.mkdir(parents=True,exist_ok=True)
    content=json.dumps({'version':1,'rows':[{'old':'invalid old row'}]})
    legacy.write_text(content,encoding='utf-8')
    dialog=ui.show_batch_analysis()
    assert dialog.rows==[]
    dialog.enqueue_files([samples[0]])
    wait_until(lambda:dialog.job is None)
    dialog.close();ui.show_batch_analysis()
    assert len(dialog.rows)==1 and legacy.read_text(encoding='utf-8')==content
    dialog.prepare_shutdown()
    assert dialog.rows==[] and legacy.read_text(encoding='utf-8')==content


def test_reopening_batch_window_keeps_rows_but_new_app_starts_empty(ui,samples):
    from tma_app.classic import ClassicMainWindow
    dialog=ui.show_batch_analysis([samples[0]])
    wait_until(lambda:dialog.job is None)
    dialog.close();assert ui.show_batch_analysis() is dialog and len(dialog.rows)==1
    other=ClassicMainWindow()
    try:assert other.show_batch_analysis().rows==[]
    finally:other.dirty=False;other.close()
    assert not (config_path().parent/'batch_history.json').exists()


def test_copy_during_export_and_new_import_does_not_change_export_snapshot(ui,samples,tmp_path,monkeypatch):
    dialog=ui.show_batch_analysis([samples[0]])
    wait_until(lambda:dialog.job is None)
    entered,release=threading.Event(),threading.Event()
    write=backend.write_xlsx
    def slow_write(rows,path):
        entered.set();release.wait(15)
        return write(rows,path)
    monkeypatch.setattr(backend,'write_xlsx',slow_write)
    target=tmp_path/'snapshot.xlsx'
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a,**k:(str(target),''))
    dialog.ExportBatchXlsx.click()
    try:
        wait_until(entered.is_set)
        assert dialog.export_job is not None and dialog.CopyBatch.isEnabled()
        assert not dialog.ExportBatchPlots.isEnabled() and not dialog.ExportBatchXlsx.isEnabled()
        dialog.CopyBatch.click();assert len(QApplication.clipboard().text().splitlines())==2
        dialog.enqueue_files([samples[-1]])
        wait_until(lambda:dialog.job is None)
        assert len(dialog.rows)==2 and dialog.export_job is not None
        assert dialog.CopyBatch.isEnabled() and not dialog.ExportBatchXlsx.isEnabled()
    finally:release.set()
    wait_until(lambda:dialog.export_job is None)
    book=load_workbook(target,data_only=True)
    assert book.active.max_row==2 and samples[0].name in book.active.cell(2,1).comment.text
    book.close()


def test_output_buttons_follow_selected_ready_rows(ui,samples,monkeypatch):
    entered,release=threading.Event(),threading.Event()
    calculate=backend.calculate_file
    def slow_second(path):
        if str(path)==str(samples[-1].resolve()):entered.set();release.wait(15)
        return calculate(path)
    monkeypatch.setattr(backend,'calculate_file',slow_second)
    dialog=ui.show_batch_analysis([samples[0],samples[-1]])
    try:
        wait_until(lambda:entered.is_set() and dialog.rows[0].get('T1') is not None)
        dialog.table.selectRow(1)
        assert all(not b.isEnabled() for b in [dialog.CopyBatch,dialog.ExportBatchXlsx,dialog.ExportBatchPlots])
        dialog.table.selectRow(0)
        assert all(b.isEnabled() for b in [dialog.CopyBatch,dialog.ExportBatchXlsx,dialog.ExportBatchPlots])
        dialog.table.clearSelection();dialog.CopyBatch.click()
        assert len(QApplication.clipboard().text().splitlines())==2
    finally:release.set()
    wait_until(lambda:dialog.job is None)


def test_export_keeps_clicked_batch_when_new_batch_arrives_during_file_picker(ui,samples,tmp_path,monkeypatch):
    dialog=ui.show_batch_analysis([samples[0]])
    wait_until(lambda:dialog.job is None)
    target=tmp_path/'clicked-batch.xlsx'
    def choose_path(*args,**kwargs):
        dialog.enqueue_files([samples[-1]])
        wait_until(lambda:dialog.job is None)
        return str(target),''
    monkeypatch.setattr(QFileDialog,'getSaveFileName',choose_path)
    dialog.ExportBatchXlsx.click()
    wait_until(lambda:dialog.export_job is None)
    book=load_workbook(target,data_only=True)
    assert book.active.max_row==2 and samples[0].name in book.active.cell(2,1).comment.text
    assert dialog.view_rows[0]['name']==samples[-1].name
    book.close()
