"""V233 parity, real XLSX/clipboard output and responsive batch UI."""
import ast
import csv
import hashlib
from io import StringIO
import logging
from pathlib import Path
import threading
import time

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication

from tma_app import batch_analysis as batch, piliang_core as core
from tma_app.classic import ClassicMainWindow
from tma_app.mass_lookup import MassMatch

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def samples():
    paths = sorted((ROOT / "测试文件").glob("*.001"))[:3]
    if not paths:
        pytest.skip("Project samples unavailable")
    return paths


def test_v14_archive_and_verbatim_dependency_closure():
    source = ROOT / "power/piliang_tma_v14.py"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == core.SOURCE_SHA256
    before = {n.name: n for n in ast.parse(source.read_text(encoding="utf-8-sig")).body
              if isinstance(n, (ast.ClassDef, ast.FunctionDef))}
    after = [n for n in ast.parse(Path(core.__file__).read_text(encoding="utf-8")).body
             if isinstance(n, (ast.ClassDef, ast.FunctionDef))]
    assert len(after) == 24
    for node in after:
        assert ast.dump(node, include_attributes=False) == ast.dump(before[node.name], include_attributes=False)
    assert "import tkinter" not in Path(core.__file__).read_text(encoding="utf-8")


def test_real_files_match_supplied_pipeline_and_remain_unchanged(samples):
    logger = logging.getLogger("test-v233")
    for path in samples:
        before = path.read_bytes()
        text, _ = core.load_text_prefer_tajm(str(path), logger)
        x, y, _ = core.extract_xy_from_file(core.parse_lines(text), logger)
        t1, t2, meta = core.detect_T2_T1_super_gradient(x, y, logger)
        cache = core.build_curve_cache(x, y)
        h = core.interpolate_curve_height(t2, cache) - core.interpolate_curve_height(t1, cache)
        actual = batch.calculate_file(path)
        assert not actual["error"]
        assert [actual[k] for k in ("T1", "T2", "H")] == pytest.approx([t1, t2, h], abs=1e-10)
        assert actual["algorithm"] == "V233" and actual["details"] == meta
        assert path.read_bytes() == before


def test_xlsx_numeric_four_columns_and_clipboard_match_display(samples, tmp_path):
    row = batch.calculate_file(samples[0])
    source = tmp_path / "weights.csv"
    source.write_text(f"样品名,M (mg)\n{samples[0].stem},0.47\n", encoding="utf-8-sig")
    original = source.read_bytes()
    row = batch.match_mass(row, str(source))
    assert row["M"] == .47
    path = tmp_path / "result.xlsx"
    batch.write_xlsx([row], path)
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    assert sheet.max_column == 4 and sheet.max_row == 2
    assert [c.value for c in sheet[1]] == list(batch.HEADERS)
    assert [c.value for c in sheet[2]] == pytest.approx([row[k] for k in batch.KEYS])
    assert all(c.data_type == "n" for c in sheet[2])
    assert sheet["C2"].number_format == "0.0000"
    assert row["file"] in sheet["A2"].comment.text and sheet.freeze_panes == "A2"
    workbook.close()
    copied = list(csv.reader(StringIO(batch.table_text([row])), delimiter="\t"))
    assert copied[1] == [batch.display_value(row, key) for key in batch.KEYS]
    assert source.read_bytes() == original
    with pytest.raises(ValueError, match="称量表"):
        batch.write_xlsx([row], source)


def test_missing_mass_is_blank_and_explicit_grams_converted(samples, tmp_path):
    row = batch.calculate_file(samples[0])
    missing = batch.match_mass(row, "")
    assert missing["M"] is None and missing["warning"]
    assert list(csv.reader(StringIO(batch.table_text([missing])), delimiter="\t"))[1][2] == ""
    source = tmp_path / "grams.csv"
    source.write_text(f"样品名,M (g)\n{samples[0].stem},0.00047\n", encoding="utf-8-sig")
    assert batch.match_mass(row, str(source))["M"] == pytest.approx(.47)


def test_plot_white_border_metadata_and_duplicate_output_names(samples, tmp_path):
    from PIL import Image
    row = batch.calculate_file(samples[0])
    paths = batch.plot_paths([row, row], tmp_path)
    assert paths[0].name == samples[0].stem + ".png" and paths[0] != paths[1]
    batch.write_plot(row, paths[0])
    with Image.open(paths[0]) as image:
        assert image.width >= 1920 + 74 and image.height >= 1080 + 74
        assert image.info["Author"] == "Power By ZzzHe"
        assert image.convert("RGB").getpixel((0, 0)) == (255, 255, 255)
    next_paths = batch.plot_paths([row], tmp_path)
    assert next_paths[0] == paths[1]


def wait_until(callback, timeout=15000):
    deadline = time.monotonic() + timeout/1000
    while not callback() and time.monotonic() < deadline:
        QTest.qWait(10)
        time.sleep(.003)
    assert callback(), "Batch operation timed out"


@pytest.fixture
def ui(monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    monkeypatch.setenv("TMA_MASS_XLSX", "")
    app = QApplication.instance() or QApplication([])
    window = ClassicMainWindow()
    window.show()
    yield window
    dialog = getattr(window, "batch_dialog", None)
    if dialog:
        dialog.cancel_processing()
        wait_until(lambda: dialog.job is None and dialog.export_job is None)
    if window.worker is not None:
        wait_until(lambda: window.worker is None)
    window.dirty = False
    window.close()
    app.processEvents()


def test_auto_on_import_multiple_files_preserves_order_and_duplicates(ui, samples):
    paths = [samples[-1], samples[0], samples[0]]
    ui.import_files(paths)
    dialog = ui.batch_dialog
    assert dialog.isVisible() and dialog.table.columnCount() == 4
    wait_until(lambda: dialog.job is None and ui.worker is None)
    assert [row["file"] for row in dialog.rows] == [str(p.resolve()) for p in paths]
    assert len({row["id"] for row in dialog.rows}) == 3
    assert all(row.get("T1") is not None for row in dialog.rows)
    assert len(ui.curves) == 3
    dialog.table.selectRow(1)
    dialog.CopyBatch.click()
    assert len(QApplication.clipboard().text().splitlines()) == 2
    assert len(QApplication.clipboard().text().splitlines()[1].split("\t")) == 4


def test_batch_failure_does_not_stop_following_files(ui, samples, tmp_path):
    bad = tmp_path / "broken.001"
    bad.write_bytes(b"broken input")
    dialog = ui.show_batch_analysis([bad, samples[0]])
    wait_until(lambda: dialog.job is None)
    assert dialog.rows[0]["error"] and dialog.rows[1]["T1"] is not None
    assert dialog.ExportBatchPlots.isEnabled() and dialog.ExportBatchXlsx.isEnabled()


def test_ui_stays_responsive_during_slow_mass_lookup_and_closes(ui, samples, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    thread_ids = []
    def slow_mass(*args, **kwargs):
        thread_ids.append(threading.get_ident())
        entered.set()
        release.wait(5)
        return MassMatch(.47, "test.csv", "mg", "matched")
    monkeypatch.setattr(batch, "lookup_record", slow_mass)
    started = time.monotonic()
    dialog = ui.show_batch_analysis([samples[0]])
    assert time.monotonic() - started < .25
    try:
        wait_until(lambda: entered.is_set() and dialog.rows[0].get("T1") is not None)
        assert dialog.rows[0].get("M") is None and dialog.isVisible()
        assert dialog.windowFlags() & Qt.WindowType.WindowCloseButtonHint
        assert dialog.close() and not dialog.isVisible()
        assert dialog.rows[0]["state"] == "已取消"
    finally:
        release.set()
    wait_until(lambda: dialog.job is None)
    assert all(t != threading.get_ident() for t in thread_ids)
    assert ui.show_batch_analysis() is dialog


def test_bottom_xlsx_and_plot_actions_really_write_outputs(ui, samples, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog
    dialog = ui.show_batch_analysis([samples[0]])
    wait_until(lambda: dialog.job is None)
    target = tmp_path / "output.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **kw: (str(target), ""))
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *a, **kw: str(tmp_path))
    dialog.ExportBatchXlsx.click()
    wait_until(lambda: dialog.export_job is None)
    assert target.exists()
    dialog.ExportBatchPlots.click()
    wait_until(lambda: dialog.export_job is None)
    assert Path(dialog.rows[0]["png"]).is_file()
    assert not dialog.rows[0].get("export_error")


def test_latest_only_all_batches_filter_and_restart(ui, samples):
    dialog = ui.show_batch_analysis([samples[0], samples[0]])
    wait_until(lambda: dialog.job is None)
    first_batch = dialog.latest_batch_id
    dialog.enqueue_files([samples[-1]])
    wait_until(lambda: dialog.job is None)
    assert len(dialog.rows) == 3 and len(dialog.view_rows) == 1
    assert dialog.table.rowCount() == 1 and dialog.table.columnCount() == 4
    dialog.all_button.click()
    assert len(dialog.view_rows) == 3
    dialog.batch_filter.setCurrentIndex(dialog.batch_filter.findData(first_batch))
    assert len(dialog.view_rows) == 2
    dialog.CopyBatch.click()
    assert len(QApplication.clipboard().text().splitlines()) == 3
    dialog.latest_button.click()
    assert len(dialog.view_rows) == 1
    from tma_app.path_settings import config_path
    assert not (config_path().parent/'batch_history.json').exists()
    stored = batch.history_snapshot(dialog.rows)
    assert len(stored) == 3 and "xs" not in stored[0] and "ys" not in stored[0]
    from tma_app.batch_dialog import BatchDialog
    reopened = BatchDialog(ui)
    assert len(reopened.rows) == 3 and len(reopened.view_rows) == 1
    reopened.all_button.click()
    assert len(reopened.view_rows) == 3
    reopened.close()


def test_popup_setting_persists_and_silent_import_still_calculates(ui, samples):
    from tma_app.path_settings import batch_popup_enabled, load_paths, save_paths
    dialog = ui.show_path_settings()
    dialog.batch_popup.setChecked(False)
    dialog.accept()
    assert not batch_popup_enabled()
    save_paths({**load_paths(), "future": "future.csv"})
    assert not batch_popup_enabled()
    ui.import_files([samples[0], samples[-1]])
    batch_dialog = ui.batch_dialog
    assert not batch_dialog.isVisible()
    wait_until(lambda: batch_dialog.job is None and ui.worker is None)
    assert len(batch_dialog.rows) == 2 and all(row.get("T1") is not None for row in batch_dialog.rows)
    assert ui.show_batch_analysis().isVisible()
    setting = ui.show_path_settings()
    assert not setting.batch_popup.isChecked()
    setting.batch_popup.setChecked(True)
    setting.accept()
    assert batch_popup_enabled()


def test_compact_result_snapshot_can_plot_and_rejects_changed_source(samples, tmp_path):
    local = tmp_path / "local.001"
    local.write_bytes(samples[0].read_bytes())
    row = batch.calculate_file(local)
    row.update(id="row", batch_id="batch", batch_label="first")
    stored = batch.history_snapshot([row])[0]
    assert "xs" not in stored
    batch.write_plot(stored, tmp_path / "history.png")
    assert (tmp_path / "history.png").is_file()
    local.write_bytes(local.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="已变化"):
        batch.write_plot(stored, tmp_path / "changed.png")


def test_queued_normal_import_adds_exactly_one_batch(ui, samples, monkeypatch):
    import tma_app.app as app_module
    read = app_module.read_dataset
    def slower_read(path):
        time.sleep(.05)
        return read(path)
    monkeypatch.setattr(app_module, "read_dataset", slower_read)
    ui.import_files([samples[0]])
    ui.import_files([samples[0], samples[-1]])
    dialog = ui.batch_dialog
    wait_until(lambda: ui.worker is None and dialog.job is None)
    assert len(ui.curves) == 3 and len(dialog.rows) == 2
    assert len({row["batch_id"] for row in dialog.rows}) == 1
